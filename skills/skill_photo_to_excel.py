"""
skill_photo_to_excel.py
-----------------------
Розпізнає таблицю на фото / скріншоті та зберігає її у структурований .xlsx.

Використовує Gemini Vision (gemini-2.5-flash) для витягування табличних даних
у JSON, а потім openpyxl для генерації Excel-файлу з форматуванням заголовків.

Тригери:
    "таблиця з фото", "перетвори в excel", "зроби таблицю",
    "розпізнай таблицю", "фото в excel", "photo to excel"

Приклад виклику:
    "зроби таблицю з C:\\screens\\receipt.jpg"
    "перетвори в excel D:\\photos\\stats.png"
"""

import os
import re
import json
import logging
from datetime import datetime

logger = logging.getLogger(__name__)

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

TRIGGERS = [
    "таблиця з фото",
    "таблицю з фото",
    "таблицю по фото",
    "таблиця по фото",
    "таблицю з фотографії",
    "таблицю зі скріншота",
    "таблицю зі скриншота",
    "перетвори в excel",
    "перетвори в ексель",
    "зроби таблицю",
    "створи таблицю",
    "створити таблицю",
    "розпізнай таблицю",
    "фото в excel",
    "фото в ексель",
    "photo to excel",
    "image to excel",
]

# Розширення зображень, які підтримує Gemini Vision
_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".gif"}

_PROMPT = """Ти експерт з розпізнавання та структурування даних із зображень.
На зображенні може бути таблиця, чек, накладна, форма, список, картка товару,
документ, скріншот, рукописні записи або інші дані, які можна подати у вигляді таблиці.

Твоє завдання — витягнути зміст і СТРОГО повернути JSON БЕЗ будь-якого тексту до або після:

{
    "title": "коротка назва набору даних (опціонально, може бути порожнім рядком)",
    "headers": ["Колонка1", "Колонка2", ...],
    "rows": [
        ["значення", "значення", ...],
        ["значення", "значення", ...]
    ]
}

АБО, якщо на зображенні є кілька окремих логічних блоків даних, поверни:

{
    "title": "загальна назва (опціонально)",
    "sheets": [
        {
            "title": "назва аркуша 1",
            "headers": ["Колонка1", "Колонка2"],
            "rows": [["...", "..."]]
        },
        {
            "title": "назва аркуша 2",
            "headers": ["Поле", "Значення"],
            "rows": [["...", "..."]]
        }
    ]
}

ПРАВИЛА НОРМАЛІЗАЦІЇ:
- Якщо на зображенні вже є таблиця, збережи її структуру.
- Якщо це список об'єктів або позицій, створи таблицю з інформативними колонками.
- Якщо це пари "поле-значення", створи таблицю з колонками ["Поле", "Значення"].
- Якщо на фото кілька незалежних секцій, таблиць або блоків, розбий їх у кілька sheet-ів.
- Якщо частина значень відсутня, використовуй порожній рядок "".
- Усі рядки повинні мати однакову кількість елементів = довжина headers.
- Заголовки мають бути короткими та зрозумілими.
- Числа, дати і коди залишай як рядки.
- Якщо змісту недостатньо для структурування, поверни {"error": "no_data"}.
- НЕ додавай ```json або інші маркери коду. Тільки чистий JSON.
"""


# Папки, де шукати фото, якщо передано лише ім'я файлу без шляху
_SEARCH_DIRS = [
    os.path.join(os.getcwd(), "uploads"),
    os.path.join(os.path.expanduser("~"), "Downloads"),
    os.path.join(os.path.expanduser("~"), "Pictures"),
    os.path.join(os.path.expanduser("~"), "Desktop"),
    os.getcwd(),
]


def _resolve_filename(name: str) -> str | None:
    """Шукає файл за іменем у відомих папках."""
    if os.path.isabs(name) and os.path.exists(name):
        return name
    for d in _SEARCH_DIRS:
        candidate = os.path.join(d, name)
        if os.path.exists(candidate):
            return candidate
    return None


def _extract_path(cmd: str) -> str | None:
    """Витягує шлях до файлу з тексту команди."""
    # Спроба 1: знайти повний шлях з диском (C:\..., D:/...)
    for ext in _IMAGE_EXTS:
        match = re.search(
            r'([A-Za-zА-Яа-яЇїІіЄєҐґ]:[\\\/][^\n\r]*?' + re.escape(ext) + r')',
            cmd,
            re.IGNORECASE,
        )
        if match:
            path = match.group(1).strip().strip('"\'')
            if os.path.exists(path):
                return path

    # Спроба 2: будь-який токен з розширенням зображення — пошук у відомих папках
    tokens = cmd.strip().split()
    for token in reversed(tokens):
        token = token.strip('"\'.,;:')
        if any(token.lower().endswith(ext) for ext in _IMAGE_EXTS):
            resolved = _resolve_filename(token)
            if resolved:
                return resolved

    # Спроба 3: весь cmd після видалення тригерів — можливо це чистий шлях/ім'я
    cleaned = cmd
    for t in TRIGGERS:
        cleaned = re.sub(re.escape(t), '', cleaned, flags=re.IGNORECASE)
    cleaned = cleaned.strip().strip('"\'').lstrip("з ").lstrip("по ").strip()
    if cleaned:
        resolved = _resolve_filename(cleaned)
        if resolved:
            return resolved

    return None


def _parse_json_response(text: str) -> dict | None:
    """Безпечно парсить JSON з відповіді LLM (з очищенням від markdown-блоків)."""
    if not text:
        return None
    # Видаляємо ```json ... ``` обгортку, якщо вона є
    text = text.strip()
    text = re.sub(r'^```(?:json)?\s*', '', text)
    text = re.sub(r'\s*```$', '', text)
    text = text.strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError as e:
        logger.warning(f"skill_photo_to_excel: JSON parse retry: {e}")

    # Іноді модель додає текст до/після — знайдемо перший JSON-об'єкт або масив
    object_start = text.find('{')
    object_end = text.rfind('}')
    array_start = text.find('[')
    array_end = text.rfind(']')

    candidates = []
    if object_start != -1 and object_end != -1 and object_end > object_start:
        candidates.append(text[object_start:object_end + 1])
    if array_start != -1 and array_end != -1 and array_end > array_start:
        candidates.append(text[array_start:array_end + 1])

    for candidate in candidates:
        try:
            return json.loads(candidate)
        except json.JSONDecodeError:
            continue

    logger.error("skill_photo_to_excel: JSON parse error: could not extract valid JSON")
    return None


def _stringify_cell(value) -> str:
    """Перетворює довільне значення у текст для Excel."""
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return str(value)
    try:
        return json.dumps(value, ensure_ascii=False)
    except TypeError:
        return str(value)


def _rows_from_records(records: list[dict], title: str = "") -> dict | None:
    """Перетворює список словників у tabular-структуру."""
    if not records:
        return None

    headers = []
    for record in records:
        if not isinstance(record, dict):
            continue
        for key in record.keys():
            key = str(key)
            if key not in headers:
                headers.append(key)

    if not headers:
        return None

    rows = []
    for record in records:
        if not isinstance(record, dict):
            continue
        rows.append([_stringify_cell(record.get(header, "")) for header in headers])

    if not rows:
        return None

    return {
        "title": title,
        "headers": headers,
        "rows": rows,
    }


def _normalize_to_tabular(data) -> dict | None:
    """Нормалізує довільний JSON у структуру {title, headers, rows}."""
    if data is None:
        return None

    if isinstance(data, dict) and data.get("error") in {"no_table", "no_data"}:
        return None

    if isinstance(data, dict):
        headers = data.get("headers")
        rows = data.get("rows")
        title = _stringify_cell(data.get("title", "")).strip()
        if isinstance(headers, list) and isinstance(rows, list) and headers:
            normalized_rows = []
            n_cols = len(headers)
            for row in rows:
                if isinstance(row, dict):
                    normalized_row = [_stringify_cell(row.get(str(header), "")) for header in headers]
                elif isinstance(row, list):
                    normalized_row = [_stringify_cell(v) for v in row[:n_cols]]
                else:
                    normalized_row = [_stringify_cell(row)]
                normalized_row += [""] * max(0, n_cols - len(normalized_row))
                normalized_rows.append(normalized_row[:n_cols])
            if normalized_rows:
                return {
                    "title": title,
                    "headers": [_stringify_cell(h) for h in headers],
                    "rows": normalized_rows,
                }

        for key in ("records", "items", "data", "entries", "rows"):
            value = data.get(key)
            if isinstance(value, list) and value:
                nested = _normalize_to_tabular(value)
                if nested:
                    if not nested.get("title"):
                        nested["title"] = title or str(key)
                    return nested

        scalar_pairs = []
        complex_records = []
        for key, value in data.items():
            if key in {"title", "headers", "rows"}:
                continue
            if isinstance(value, dict):
                flattened = {"Поле": str(key)}
                for nested_key, nested_value in value.items():
                    flattened[str(nested_key)] = _stringify_cell(nested_value)
                complex_records.append(flattened)
            elif isinstance(value, list):
                if value and all(isinstance(item, dict) for item in value):
                    nested = _rows_from_records(value, title or str(key))
                    if nested:
                        return nested
                else:
                    scalar_pairs.append([str(key), _stringify_cell(value)])
            else:
                scalar_pairs.append([str(key), _stringify_cell(value)])

        if complex_records:
            nested = _rows_from_records(complex_records, title)
            if nested:
                return nested

        if scalar_pairs:
            return {
                "title": title,
                "headers": ["Поле", "Значення"],
                "rows": scalar_pairs,
            }

        return None

    if isinstance(data, list):
        if not data:
            return None
        if all(isinstance(item, dict) for item in data):
            return _rows_from_records(data)
        if all(isinstance(item, list) for item in data):
            n_cols = max(len(item) for item in data)
            headers = [f"Колонка {idx}" for idx in range(1, n_cols + 1)]
            rows = []
            for item in data:
                row = [_stringify_cell(value) for value in item[:n_cols]]
                row += [""] * max(0, n_cols - len(row))
                rows.append(row)
            return {
                "title": "",
                "headers": headers,
                "rows": rows,
            }
        return {
            "title": "",
            "headers": ["Значення"],
            "rows": [[_stringify_cell(item)] for item in data],
        }

    return {
        "title": "",
        "headers": ["Значення"],
        "rows": [[_stringify_cell(data)]],
    }


def _normalize_to_workbook(data) -> dict | None:
    """Нормалізує довільний JSON у workbook-структуру з одним або кількома аркушами."""
    if data is None:
        return None

    if isinstance(data, dict) and isinstance(data.get("sheets"), list) and data.get("sheets"):
        workbook_title = _stringify_cell(data.get("title", "")).strip()
        sheets = []
        for index, sheet_data in enumerate(data.get("sheets", []), start=1):
            normalized_sheet = _normalize_to_tabular(sheet_data)
            if not normalized_sheet:
                continue
            if not normalized_sheet.get("title"):
                normalized_sheet["title"] = f"Аркуш {index}"
            sheets.append(normalized_sheet)
        if sheets:
            return {
                "title": workbook_title,
                "sheets": sheets,
            }
        return None

    normalized = _normalize_to_tabular(data)
    if not normalized:
        return None
    return {
        "title": _stringify_cell(normalized.get("title", "")).strip(),
        "sheets": [normalized],
    }


def _sheet_name(title: str, index: int) -> str:
    """Повертає безпечну назву аркуша Excel."""
    name = (title or f"Аркуш {index}").strip()
    name = re.sub(r'[\\/*?:\[\]]', "_", name)
    return name[:31] or f"Аркуш {index}"


def _fill_sheet(ws, data: dict) -> int:
    """Заповнює аркуш даними та повертає кількість рядків."""
    headers = data.get("headers") or []
    rows = data.get("rows") or []
    title = (data.get("title") or "").strip()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    start_row = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
        if headers:
            ws.merge_cells(start_row=1, start_column=1,
                           end_row=1, end_column=max(1, len(headers)))
        start_row = 3

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=str(h))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    n_cols = len(headers) if headers else (max(len(r) for r in rows) if rows else 0)
    for r_idx, row in enumerate(rows, start=start_row + 1):
        if not isinstance(row, list):
            continue
        normalized = list(row[:n_cols]) + [""] * max(0, n_cols - len(row))
        for c_idx, val in enumerate(normalized, start=1):
            cell_value = val
            if isinstance(val, str):
                stripped = val.strip().replace(",", ".")
                if re.fullmatch(r'-?\d+(?:\.\d+)?', stripped):
                    try:
                        cell_value = float(stripped) if "." in stripped else int(stripped)
                    except ValueError:
                        cell_value = val
            cell = ws.cell(row=r_idx, column=c_idx, value=cell_value)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col_idx in range(1, n_cols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 10
        for r_idx in range(start_row, start_row + 1 + len(rows)):
            v = ws.cell(row=r_idx, column=col_idx).value
            if v is not None:
                max_len = max(max_len, min(50, len(str(v)) + 2))
        ws.column_dimensions[col_letter].width = max_len

    if title:
        ws.row_dimensions[start_row].height = 28
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return len(rows)


def _write_xlsx(data: dict, dest_path: str) -> int:
    """Записує workbook-структуру у .xlsx. Повертає загальну кількість рядків."""
    wb = Workbook()
    sheets = data.get("sheets") or []
    total_rows = 0

    if not sheets:
        wb.active.title = "Розпізнано"
        _fill_sheet(wb.active, {"title": data.get("title", ""), "headers": [], "rows": []})
    else:
        first_sheet = sheets[0]
        ws = wb.active
        ws.title = _sheet_name(first_sheet.get("title", "Розпізнано"), 1)
        total_rows += _fill_sheet(ws, first_sheet)

        for index, sheet_data in enumerate(sheets[1:], start=2):
            ws = wb.create_sheet(title=_sheet_name(sheet_data.get("title", ""), index))
            total_rows += _fill_sheet(ws, sheet_data)

    wb.save(dest_path)
    return total_rows


def execute(cmd: str, jarvis) -> str:
    if Workbook is None:
        return "Потрібна бібліотека openpyxl. Встановіть: pip install openpyxl"
    if PILImage is None:
        return "Потрібна бібліотека Pillow."

    if not getattr(jarvis.llm, "gemini_client", None):
        return "Для розпізнавання таблиць потрібен Gemini. Активуйте провайдер Gemini."

    image_path = _extract_path(cmd)
    if not image_path:
        return ("Не можу знайти файл зображення в команді. "
                "Скажіть, наприклад: 'зроби таблицю з C:\\photos\\table.jpg'")

    ext = os.path.splitext(image_path)[1].lower()
    if ext not in _IMAGE_EXTS:
        return f"Формат {ext} не підтримується. Використовуйте JPG, PNG, BMP, WEBP."

    jarvis.audio.speak(f"Розпізнаю таблицю на {os.path.basename(image_path)}...")

    try:
        img = PILImage.open(image_path)
    except Exception as e:
        return f"Не вдалося відкрити зображення: {e}"

    try:
        res = jarvis.llm.gemini_client.models.generate_content(
            model="gemini-2.5-flash",
            contents=[img, _PROMPT],
        )
        raw_text = (res.text or "").strip()
    except Exception as e:
        logger.error(f"skill_photo_to_excel: Gemini error: {e}")
        return f"Помилка Vision API: {e}"

    raw_data = _parse_json_response(raw_text)
    if raw_data is None:
        logger.warning(f"skill_photo_to_excel: невалідний JSON: {raw_text[:300]}")
        return "Не вдалося розпізнати структуру таблиці. Спробуйте чіткіше фото."

    data = _normalize_to_workbook(raw_data)
    if data is None or not data.get("sheets"):
        return "Не вдалося структурувати дані з фото у таблицю."

    # Куди зберігати — окрема папка Converted_Excel у корені проєкту
    out_dir = os.path.join(os.getcwd(), "Converted_Excel")
    try:
        os.makedirs(out_dir, exist_ok=True)
    except Exception as e:
        logger.error(f"skill_photo_to_excel: не вдалося створити {out_dir}: {e}")
        out_dir = os.path.dirname(os.path.abspath(image_path)) or os.getcwd()
    base_name = os.path.splitext(os.path.basename(image_path))[0]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest_path = os.path.join(out_dir, f"{base_name}_table_{timestamp}.xlsx")

    try:
        n_rows = _write_xlsx(data, dest_path)
    except Exception as e:
        logger.error(f"skill_photo_to_excel: write error: {e}", exc_info=True)
        return f"Помилка запису Excel: {e}"

    n_sheets = len(data.get("sheets", []))
    max_cols = max((len(sheet.get("headers", [])) for sheet in data.get("sheets", [])), default=0)
    logger.info(f"skill_photo_to_excel: створено {dest_path} ({n_rows} рядків, {max_cols} колонок, {n_sheets} аркушів)")

    # Спроба автовідкриття на Windows
    try:
        if os.name == "nt":
            os.startfile(dest_path)  # noqa: S606  # nosec
    except Exception:
        pass

    return (f"Готово. Створено таблицю {n_rows} рядків, до {max_cols} колонок і {n_sheets} аркушів. "
            f"Файл збережено: {dest_path}")
