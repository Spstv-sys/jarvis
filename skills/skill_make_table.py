"""
skill_make_table.py
-------------------
Створює структуровану Excel-таблицю з документа (PDF, DOCX, TXT, CSV, MD)
або зображення, надсилаючи файл саме до GPT (OpenAI):

- Текстові формати: локально витягуємо текст і передаємо у GPT як промпт.
- Зображення: надсилаються напряму через Vision API (base64) у gpt-4o*.

Формат відповіді — той самий JSON-схема, що й у skill_photo_to_excel,
тож переvикористовуємо `_normalize_to_workbook` і `_write_xlsx`.

Тригери (навмисно відрізняються від skill_photo_to_excel,
щоб уникнути конфліктів маршрутизації):

    "створи таблицю з документа", "створи таблицю з документу",
    "створи таблицю з файлу",     "створи таблицю з pdf",
    "зроби таблицю з документа",  "зроби таблицю з файлу",
    "таблиця з документа",        "таблиця з документу",
    "таблицю з документа",        "таблицю з документу",
    "таблиця з pdf",              "таблицю з pdf",
    "перетвори документ в excel", "перетвори документ в ексель",
    "перетвори pdf в excel",      "перетвори pdf в ексель",
    "gpt таблиця",                "gpt таблицю"

Приклад:
    "створи таблицю з документа D:\\docs\\order.pdf"
    "перетвори pdf в excel C:\\reports\\summary.pdf"
"""

import logging
import os
import re
from datetime import datetime

logger = logging.getLogger(__name__)

# Перevикористовуємо хелпери з фото-скіла, щоб не дублювати код
from skills.skill_photo_to_excel import (  # type: ignore
    _normalize_to_workbook,
    _parse_json_response,
    _write_xlsx,
)

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx  # python-docx
except ImportError:
    docx = None


TRIGGERS = [
    "створи таблицю з документа",
    "створи таблицю з документу",
    "створи таблицю з файлу",
    "створи таблицю з pdf",
    "зроби таблицю з документа",
    "зроби таблицю з документу",
    "зроби таблицю з файлу",
    "зроби таблицю з pdf",
    "таблиця з документа",
    "таблиця з документу",
    "таблицю з документа",
    "таблицю з документу",
    "таблиця з pdf",
    "таблицю з pdf",
    "перетвори документ в excel",
    "перетвори документ в ексель",
    "перетвори pdf в excel",
    "перетвори pdf в ексель",
    "gpt таблиця",
    "gpt таблицю",
]

_DOC_EXTS = {".pdf", ".docx", ".txt", ".md", ".csv", ".log", ".json"}
_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".gif"}
_ALL_EXTS = _DOC_EXTS | _IMAGE_EXTS

_MAX_TEXT_CHARS = 60_000  # обрізаємо дуже великі документи, щоб не вилетіти за context window

_SEARCH_DIRS = [
    os.path.join(os.getcwd(), "uploads"),
    os.path.join(os.path.expanduser("~"), "Downloads"),
    os.path.join(os.path.expanduser("~"), "Documents"),
    os.path.join(os.path.expanduser("~"), "Desktop"),
    os.getcwd(),
]

_PROMPT_TEMPLATE = """Ти експерт з структурування даних.
Користувач передає вміст документа і просить побудувати таблицю.

Поверни СТРОГО JSON БЕЗ будь-яких пояснень, маркерів коду або тексту:

{{
    "title": "коротка назва таблиці (опціонально)",
    "headers": ["Колонка1", "Колонка2", ...],
    "rows": [
        ["значення", "значення", ...]
    ]
}}

АБО, якщо в документі логічно кілька таблиць/блоків:

{{
    "title": "загальна назва",
    "sheets": [
        {{ "title": "...", "headers": [...], "rows": [[...]] }}
    ]
}}

ПРАВИЛА:
- Заголовки короткі та інформативні.
- Усі рядки однакової довжини = len(headers); відсутні значення = "".
- Числа/дати/коди залишай як рядки.
- Якщо неможливо побудувати таблицю — поверни {{"error": "no_data"}}.
- Жодних ```json блоків. Тільки чистий JSON.

ДОДАТКОВА ВКАЗІВКА КОРИСТУВАЧА: {user_hint}

ВМІСТ ДОКУМЕНТА:
---
{content}
---
"""


def _resolve_filename(name: str) -> str | None:
    if os.path.isabs(name) and os.path.exists(name):
        return name
    for d in _SEARCH_DIRS:
        candidate = os.path.join(d, name)
        if os.path.exists(candidate):
            return candidate
    return None


def _extract_path(cmd: str) -> str | None:
    # 1) Абсолютний шлях з диском
    for ext in _ALL_EXTS:
        m = re.search(
            r'([A-Za-zА-Яа-яЇїІіЄєҐґ]:[\\\/][^\n\r]*?' + re.escape(ext) + r')',
            cmd, re.IGNORECASE,
        )
        if m:
            path = m.group(1).strip().strip('"\'')
            if os.path.exists(path):
                return path

    # 2) Останній токен з відомим розширенням
    for token in reversed(cmd.split()):
        token = token.strip('"\'.,;:')
        if any(token.lower().endswith(ext) for ext in _ALL_EXTS):
            resolved = _resolve_filename(token)
            if resolved:
                return resolved

    # 3) Залишок після видалення тригерів
    cleaned = cmd
    for t in TRIGGERS:
        cleaned = re.sub(re.escape(t), "", cleaned, flags=re.IGNORECASE)
    cleaned = cleaned.strip().strip('"\'').lstrip("з ").lstrip("із ").lstrip("по ").strip()
    if cleaned:
        return _resolve_filename(cleaned)
    return None


def _extract_text(path: str) -> str | None:
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            if fitz is None:
                return None
            text_parts: list[str] = []
            with fitz.open(path) as doc:
                for page in doc:
                    text_parts.append(page.get_text())
            return "\n".join(text_parts).strip() or None

        if ext == ".docx":
            if docx is None:
                return None
            d = docx.Document(path)
            paragraphs = [p.text for p in d.paragraphs if p.text]
            # Прості таблиці теж витягуємо
            for tbl in d.tables:
                for row in tbl.rows:
                    paragraphs.append(" | ".join(c.text for c in row.cells))
            return "\n".join(paragraphs).strip() or None

        if ext in {".txt", ".md", ".csv", ".log", ".json"}:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read().strip() or None
    except Exception as e:
        logger.error(f"skill_make_table: text extract error from {path}: {e}")
        return None

    return None


def _strip_triggers(cmd: str) -> str:
    out = cmd
    for t in TRIGGERS:
        out = re.sub(re.escape(t), "", out, flags=re.IGNORECASE)
    # також прибираємо шлях, якщо він залишився
    out = re.sub(r'[A-Za-zА-Яа-яЇїІіЄєҐґ]:[\\\/]\S+', "", out)
    out = re.sub(r'\S+\.(pdf|docx|txt|md|csv|log|json|jpe?g|png|bmp|webp|gif)\b',
                 "", out, flags=re.IGNORECASE)
    return out.strip().strip('"\'').strip("-,;:") or "Побудуй найбільш інформативну таблицю."


def execute(cmd: str, jarvis) -> str:
    openai_provider = jarvis.llm._providers.get("openai")
    if openai_provider is None or not openai_provider.available:
        return ("Для цієї функції потрібен OpenAI-провайдер. "
                "Встановіть OPENAI_API_KEY і перезапустіть.")

    path = _extract_path(cmd)
    if not path:
        return ("Не можу знайти файл у команді. "
                "Скажіть, наприклад: 'створи таблицю з документа D:\\\\docs\\\\order.pdf'")

    ext = os.path.splitext(path)[1].lower()
    if ext not in _ALL_EXTS:
        return f"Формат {ext} не підтримується. Дозволено: PDF, DOCX, TXT, MD, CSV, JSON або зображення."

    user_hint = _strip_triggers(cmd)
    jarvis.audio.speak(f"Надсилаю {os.path.basename(path)} до GPT для побудови таблиці...")

    raw_text: str | None = None

    if ext in _IMAGE_EXTS:
        # Vision: файл йде напряму до GPT
        prompt = _PROMPT_TEMPLATE.format(user_hint=user_hint, content="(надано зображення)")
        raw_text = openai_provider.generate_with_image(prompt, path)
    else:
        content = _extract_text(path)
        if not content:
            return (f"Не вдалося витягнути текст з {os.path.basename(path)}. "
                    "Можливо, відсутні бібліотеки PyMuPDF/python-docx або файл порожній.")
        if len(content) > _MAX_TEXT_CHARS:
            logger.info(f"skill_make_table: обрізаю вміст {len(content)} → {_MAX_TEXT_CHARS} символів")
            content = content[:_MAX_TEXT_CHARS] + "\n...[обрізано]..."
        prompt = _PROMPT_TEMPLATE.format(user_hint=user_hint, content=content)
        raw_text = openai_provider.generate(prompt)

    if not raw_text:
        return "GPT не повернув відповіді. Спробуйте пізніше або перевірте ліміти/ключ."

    parsed = _parse_json_response(raw_text)
    if parsed is None:
        logger.warning(f"skill_make_table: невалідний JSON: {raw_text[:300]}")
        return "Не вдалося розпарсити відповідь GPT як JSON. Спробуйте ще раз."

    if isinstance(parsed, dict) and parsed.get("error") in {"no_data", "no_table"}:
        return "GPT вирішив, що в документі немає даних, придатних для таблиці."

    workbook = _normalize_to_workbook(parsed)
    if workbook is None or not workbook.get("sheets"):
        return "Не вдалося структурувати дані у таблицю."

    out_dir = os.path.join(os.getcwd(), "Converted_Excel")
    try:
        os.makedirs(out_dir, exist_ok=True)
    except Exception as e:
        logger.error(f"skill_make_table: не вдалося створити {out_dir}: {e}")
        out_dir = os.path.dirname(os.path.abspath(path)) or os.getcwd()

    base = os.path.splitext(os.path.basename(path))[0]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = os.path.join(out_dir, f"{base}_gpt_table_{ts}.xlsx")

    try:
        n_rows = _write_xlsx(workbook, dest)
    except Exception as e:
        logger.error(f"skill_make_table: write error: {e}", exc_info=True)
        return f"Помилка запису Excel: {e}"

    n_sheets = len(workbook.get("sheets", []))
    max_cols = max((len(s.get("headers", [])) for s in workbook.get("sheets", [])), default=0)
    logger.info(f"skill_make_table: створено {dest} ({n_rows} рядків, {max_cols} колонок, {n_sheets} аркушів)")

    try:
        if os.name == "nt":
            os.startfile(dest)  # noqa: S606  # nosec
    except Exception:
        pass

    return (f"Готово. GPT побудував таблицю: {n_rows} рядків, до {max_cols} колонок, "
            f"{n_sheets} аркуш(ів). Файл: {dest}")
